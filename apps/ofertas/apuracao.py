"""Arquivo de apuração pra enviar à indústria (docx, pendência aberta desde
13/09/26) — base = relatório de vendas por item, linha a linha (não
agregado), + investimento que o fabricante deve pagar. 2 fórmulas
diferentes, dependendo da mecânica:

- Ofertas de fabricante (Kenvue/Principia/Botica/Procter, `grupo=oferta`):
  `montar_apuracao_industria` -- valor de rebaixa por EAN (`RebaixaProduto`,
  o fabricante manda a tabela; 1ª implementação real: Procter Semana do
  Cliente).
- Leve 3 Pague 2: `montar_apuracao_leve3` -- fórmula própria automática já
  validada (ciclos × custo), não depende de tabela de rebaixa nenhuma.

Cestões/Marketing/Kimberly/Supra Corp ainda não têm fórmula definida (ver
README, seção Pendências) -- não têm função de apuração ainda.
"""
from __future__ import annotations

from decimal import Decimal

import pandas as pd
from openpyxl.styles import Border, Font, Side

from apps.produtos.models import Produto, RebaixaProduto

from .erp import tag_sem_prefixo
from .models import Lancamento

ZERO = Decimal('0')

_COLUNAS_FABRICANTE = {
    'bandeira': 'Bandeira', 'data': 'Data', 'ean': 'EAN',
    'produto': 'Produto', 'itens': 'Itens', 'venda': 'Venda',
    'valor_rebaixa_unitario': 'Valor da Rebaixa (R$/un.)',
    'investimento': 'Investimento (R$)',
}
_COLUNAS_LEVE3 = {
    'bandeira': 'Bandeira', 'data': 'Data', 'ano_mes': 'Ano-mês',
    'produto': 'Produto', 'itens': 'Itens', 'venda': 'Venda',
    'ciclos': 'Ciclos', 'custo_unitario': 'Custo Unitário (R$)',
    'investimento': 'Investimento (R$)',
}

# Colunas somadas na linha de TOTAL do rodapé -- Custo/Lucro/Desconto
# nunca entram nesse arquivo (pedido do Gabriel 22/09, "não faz sentido
# mostrar nossa margem pra indústria" -- não são a base de nenhuma das 2
# fórmulas de investimento, só "Custo Unitário" do Leve3 é, e esse fica
# de fora da soma por ser um valor unitário, não aditivo).
COLUNAS_SOMAVEIS = {'Itens', 'Venda', 'Ciclos', 'Investimento (R$)'}


def _slug(texto: str) -> str:
    """'OFERTAS PROCTER SEMANA DO CLIENTE' -> 'ofertas_procter_semana_do_cliente'."""
    return texto.strip().lower().replace(' ', '_')


def nome_arquivo_apuracao(nome_oferta: str, ano_mes: str | None) -> str:
    """Convenção pedida pelo Gabriel (22/09): "apuracao_oferta_mes" --
    `nome_oferta` já vem pronto pra usar (campanha, ou fabricante no caso
    do Leve3, que precisa de 1 arquivo por fabricante). Sem mês (`ano_mes`
    omitido) cai pra "todos_os_meses", pra não perder a distinção de quando
    o arquivo é um recorte único ou o histórico inteiro."""
    return f'apuracao_{_slug(nome_oferta)}_{ano_mes or "todos_os_meses"}.xlsx'


def fabricantes_leve3(ano_mes: str | None = None) -> list[str]:
    """Fabricantes com lançamento no Leve3 (opcionalmente só 1 mês) --
    Gabriel pediu o arquivo de apuração separado por fabricante, não 1 só
    com tudo junto."""
    queryset = Lancamento.objects.filter(mecanica=Lancamento.LEVE3).exclude(fabricante='')
    if ano_mes:
        queryset = queryset.filter(ano_mes=ano_mes)
    return sorted(queryset.values_list('fabricante', flat=True).distinct())


def gerar_dataframe_apuracao(
    mecanica: str, campanha: str | None = None, ano_mes: str | None = None, fabricante: str | None = None,
):
    """Monta os dados (`montar_apuracao_leve3`/`montar_apuracao_industria`,
    conforme a mecânica) e devolve `(dataframe_pronto_pra_excel, dados)` --
    usado tanto pelo management command quanto pelo botão de download no
    painel, pra não duplicar a lógica de renomear coluna/arredondar em 2
    lugares. `fabricante` só vale pro Leve3 (Kenvue/Principia/Botica/
    Procter já são 1 fabricante só por mecânica, não precisa filtrar)."""
    if mecanica == Lancamento.LEVE3:
        dados = montar_apuracao_leve3(ano_mes=ano_mes, fabricante=fabricante)
        colunas = _COLUNAS_LEVE3
    else:
        if not campanha:
            raise ValueError('campanha é obrigatória pra essa mecânica (só o Leve3 dispensa).')
        dados = montar_apuracao_industria(mecanica, campanha, ano_mes=ano_mes)
        colunas = _COLUNAS_FABRICANTE

    df = pd.DataFrame(dados['linhas'])
    if not df.empty:
        # Arredonda só aqui, pra exibir -- `dados['total_investimento']`
        # já foi somado em precisão cheia antes disso (ver comentário em
        # `montar_apuracao_leve3` sobre a diferença de 7 centavos achada
        # ao conferir contra o painel).
        df['investimento'] = df['investimento'].astype(float).round(2)
    df = df.rename(columns=colunas)
    return df, dados


def periodo_texto(dados: dict) -> str:
    """Texto de período pro cabeçalho do arquivo de apuração -- 'Período:
    dd/mm/aaaa a dd/mm/aaaa', sempre a partir da DATA REAL das linhas
    exportadas (a data em que a ação/promoção de fato vendeu), nunca o
    mês-calendário inteiro do filtro (correção 22/09/26: "o cabeçalho tem
    que ser a data da ação" -- 1ª versão usava dia 1 ao último dia do mês
    quando `?mes=` vinha preenchido, mas a ação pode ter rodado só numa
    semana daquele mês, ex. Leve3/Procter)."""
    datas = [l['data'] for l in dados['linhas'] if l.get('data')]
    if datas:
        return f'Período: {min(datas).strftime("%d/%m/%Y")} a {max(datas).strftime("%d/%m/%Y")}'
    return 'Período: sem data (lançamento sem dia importado)'


def resumo_apuracao_leve3(ano_mes: str | None = None) -> list[dict]:
    """Resumo por fabricante pra mostrar na seção "Apuração para a
    indústria" sem precisar abrir cada Excel -- pedido do Gabriel
    (22/09/26, "como ele está visível na nossa ferramenta"): 1 linha por
    fabricante com nome, período real, nº de linhas e investimento total,
    ordenado do maior pro menor. Reaproveita `montar_apuracao_leve3` (1
    chamada por fabricante -- aceitável, são só ~9 hoje) em vez de duplicar
    a lógica de cálculo."""
    resumo = []
    for fabricante in fabricantes_leve3(ano_mes=ano_mes):
        dados = montar_apuracao_leve3(ano_mes=ano_mes, fabricante=fabricante)
        if not dados['linhas']:
            continue
        resumo.append({
            'nome': fabricante,
            'periodo': periodo_texto(dados),
            'linhas': len(dados['linhas']),
            'investimento': dados['total_investimento'],
        })
    resumo.sort(key=lambda r: r['investimento'], reverse=True)
    return resumo


def resumo_apuracao_industria(mecanica: str, campanhas: list[str], ano_mes: str | None = None) -> list[dict]:
    """Mesma ideia do `resumo_apuracao_leve3`, pras mecânicas de
    fabricante (Kenvue/Principia/Botica/Procter) -- 1 linha por campanha
    que já tem rebaixa cadastrada (`campanhas`, vem de `RebaixaProduto`)."""
    resumo = []
    for campanha in campanhas:
        dados = montar_apuracao_industria(mecanica, campanha, ano_mes=ano_mes)
        if not dados['linhas']:
            continue
        resumo.append({
            'nome': campanha,
            'periodo': periodo_texto(dados),
            'linhas': len(dados['linhas']),
            'investimento': dados['total_investimento'],
        })
    resumo.sort(key=lambda r: r['investimento'], reverse=True)
    return resumo


def escrever_excel_apuracao(df, dados: dict, nome_oferta: str, destino) -> None:
    """Escreve o .xlsx de apuração com o mesmo cabeçalho de 2 linhas do
    relatório "Análise de Venda por Item" do ERP (título + "Período: ...")
    antes da tabela -- pedido do Gabriel (22/09), com foto do relatório
    de referência -- e uma linha de TOTAL no rodapé (pedido 22/09,
    reformulação: "não tem total nenhum... a indústria teria que somar
    2 mil linhas na mão"). `destino` é um caminho (`Path`/str, usado pelo
    management command) ou um objeto tipo-arquivo (`HttpResponse`, usado
    pelo botão de download no painel) -- os dois funcionam igual com
    `pd.ExcelWriter`."""
    with pd.ExcelWriter(destino, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Apuração', startrow=2)
        ws = writer.sheets['Apuração']
        ws['A1'] = f'Apuração {nome_oferta}'
        ws['A1'].font = Font(bold=True)
        ws['A2'] = periodo_texto(dados)

        if not df.empty:
            linha_total = 4 + len(df)  # 1=título, 2=período, 3=cabeçalho, 4..=dados
            borda_topo = Border(top=Side(style='thin'))
            ws.cell(row=linha_total, column=1, value='TOTAL').font = Font(bold=True)
            for coluna, indice in zip(df.columns, range(1, len(df.columns) + 1)):
                celula = ws.cell(row=linha_total, column=indice)
                if coluna == 'Investimento (R$)':
                    # NÃO soma a coluna do DataFrame aqui -- ela já veio
                    # arredondada linha a linha (pra exibição), somar as
                    # linhas arredondadas dá um total ligeiramente diferente
                    # do valor certo (mesmo bug de precisão já corrigido
                    # antes, achado de novo ao testar esta linha de TOTAL:
                    # R$7.698,42 vs R$7.698,48 reais). `dados['total_investimento']`
                    # já vem somado em precisão cheia, arredonda só aqui.
                    celula.value = round(float(dados['total_investimento']), 2)
                    celula.font = Font(bold=True)
                elif coluna in COLUNAS_SOMAVEIS:
                    celula.value = float(df[coluna].sum())
                    celula.font = Font(bold=True)
                celula.border = borda_topo


def montar_apuracao_industria(mecanica: str, campanha: str, ano_mes: str | None = None) -> dict:
    """Devolve as linhas de venda da campanha (grupo=oferta) + rebaixa/
    investimento por linha, prontas pra exportar. Resolve o EAN de cada
    produto pelo cadastro (`Produto.descricao` -> `codigo_barras`) e cruza
    com `RebaixaProduto`; quando o cadastro não tem EAN pro produto (achado
    real: acontece, ver commit), cai pro 2º critério — nome do produto
    batendo exato com `RebaixaProduto.produto_descricao`. `ano_mes`
    (opcional, "AAAA-MM") filtra pra 1 mês só; omitido = todos."""
    queryset = Lancamento.objects.filter(mecanica=mecanica, grupo=Lancamento.GRUPO_OFERTA)
    if ano_mes:
        queryset = queryset.filter(ano_mes=ano_mes)
    linhas = list(
        queryset.select_related('loja').values(
            'loja__bandeira', 'data', 'ano_mes',
            'produto_descricao', 'tag_origem', 'itens', 'venda',
        )
    )
    linhas = [l for l in linhas if tag_sem_prefixo(l['tag_origem']) == campanha]

    descricoes = {l['produto_descricao'] for l in linhas}
    ean_por_descricao = dict(
        Produto.objects.filter(descricao__in=descricoes).values_list('descricao', 'codigo_barras')
    )

    rebaixas = list(RebaixaProduto.objects.filter(mecanica=mecanica, campanha=campanha))
    rebaixa_por_ean = {r.ean: r.valor_rebaixa for r in rebaixas if r.ean}
    rebaixa_por_nome = {r.produto_descricao: r.valor_rebaixa for r in rebaixas}

    resultado = []
    sem_rebaixa = set()
    for l in linhas:
        descricao = l['produto_descricao']
        ean = ean_por_descricao.get(descricao, '')
        valor_rebaixa = rebaixa_por_ean.get(ean) if ean else None
        if valor_rebaixa is None:
            valor_rebaixa = rebaixa_por_nome.get(descricao)
        if valor_rebaixa is None:
            sem_rebaixa.add(descricao)
            valor_rebaixa = ZERO

        itens = l['itens'] or ZERO
        resultado.append({
            'bandeira': l['loja__bandeira'],
            'data': l['data'],
            'ean': ean,
            'produto': descricao,
            'itens': itens,
            'venda': l['venda'] or ZERO,
            'valor_rebaixa_unitario': valor_rebaixa,
            # Sem quantize aqui pelo mesmo motivo do Leve3 (ver
            # montar_apuracao_leve3) -- soma em precisão cheia, arredondar
            # só na exportação.
            'investimento': itens * valor_rebaixa,
        })

    return {
        'linhas': resultado,
        'total_investimento': sum((r['investimento'] for r in resultado), ZERO),
        'total_itens': sum((r['itens'] for r in resultado), ZERO),
        'produtos_sem_rebaixa': sorted(sem_rebaixa),
    }


def montar_apuracao_leve3(ano_mes: str | None = None, fabricante: str | None = None) -> dict:
    """Leve 3 Pague 2 tem fórmula própria automática (docx, seção 5.1/6) --
    não depende de tabela de rebaixa nenhuma, diferente das ofertas de
    fabricante. Mesmo cálculo linha a linha do `calcular_leve3`
    (`services.py`), só que devolvendo cada linha pronta pra exportar em
    vez de já agregada: ciclos = itens // 3; investimento = ciclos ×
    (custo da linha / itens da linha). `ano_mes` (opcional, "AAAA-MM")
    filtra pra 1 mês só; omitido = todos. `fabricante` (opcional, ex.
    "EMS") -- Gabriel pediu 1 arquivo por fabricante, não 1 só com todos
    juntos (docx cobre genéricos de vários laboratórios no mesmo Leve3)."""
    queryset = Lancamento.objects.filter(mecanica=Lancamento.LEVE3)
    if ano_mes:
        queryset = queryset.filter(ano_mes=ano_mes)
    if fabricante:
        queryset = queryset.filter(fabricante=fabricante)
    linhas = queryset.select_related('loja').values(
        'loja__bandeira', 'data', 'ano_mes',
        'produto_descricao', 'itens', 'venda', 'custo',
    )

    resultado = []
    for l in linhas:
        itens = l['itens'] or ZERO
        custo = l['custo'] or ZERO
        custo_unitario = (custo / itens) if itens else ZERO
        ciclos = itens // 3
        # Sem arredondar aqui -- arredondar cada linha pra 2 casas e DEPOIS
        # somar dá um total diferente da soma em precisão cheia (achado
        # real: R$50.646,77 vs R$50.646,84 do painel, 3.286 linhas, 7
        # centavos de diferença acumulada). O painel (`calcular_leve3`)
        # nunca arredonda linha a linha -- fica igual aqui, o
        # arredondamento pra 2 casas só acontece na exportação (cosmético,
        # não usado pra somar `total_investimento` abaixo).
        investimento = ciclos * custo_unitario
        resultado.append({
            'bandeira': l['loja__bandeira'],
            'data': l['data'],
            'ano_mes': l['ano_mes'],
            'produto': l['produto_descricao'],
            'itens': itens,
            'venda': l['venda'] or ZERO,
            'ciclos': ciclos,
            'custo_unitario': custo_unitario.quantize(Decimal('0.0001')),
            'investimento': investimento,
        })

    return {
        'linhas': resultado,
        'total_investimento': sum((r['investimento'] for r in resultado), ZERO),
        'total_itens': sum((r['itens'] for r in resultado), ZERO),
    }
